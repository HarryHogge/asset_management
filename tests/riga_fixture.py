"""A made-up Riga workbook in the layout report 10 describes. Every serial,
note, initial and value here is invented; no real Riga data is used.

What it reproduces from report 10:
- 10 equipment sheets with a merged title in row 1, a merged color legend in
  row 2 and the headers in row 3; PPC with its headers in row 2;
- header names, case and column order as report 10 lists them per sheet
  (Status in column E on most sheets, D on TRC, B on GCS; STATUS on EPS and
  GCS; NOTES on VXE, INFO on PPC, INFO: on GCS; Insurance and SIGN columns);
- a side table to the right (2.5 VTOL columns W to Y, TRC's count table in
  M and N, ENG's marker columns L and M) and one below with no serial
  (2.5 VTOL rows 21 to 25);
- an Excel date, dd.mm.yyyy and dd.mm.yy text dates, a date that cannot be
  read, and a date written inside the notes;
- a serial stored as a number (VXE);
- a header in different case from the mapping (STATUS on VXE);
- an unknown Status value (NFC ISSUES on EPS);
- a formula error (#REF!) in a mapped column (MK3's Location) and in the
  Equip. Rediness columns, which are not read;
- a missing mapped header (MK2 calls its owner column "Owner");
- the three ignored sheets, and one sheet the mapping does not know;
- an engine sheet with a mix of conditions.
"""
from datetime import datetime

from openpyxl import Workbook

# Initials in the SIGN columns. The check command must never print them.
SIGN = 'QZX'
# Text in the notes. The check command must never print it.
SECRET_NOTE = 'Secret note text 7781'

TITLE_LEGEND = ('FLIGHT READY', 'NOT FLIGHT READY', 'Insurance expired')

AIRCRAFT_HEADERS = ['S/N', 'Configuration', 'Frequency', 'Insurance status', 'Status', 'Location',
                    'Asset owner', 'INFO UPDATE', 'SIGN', 'Actions/Defects', 'NC #',
                    'Equip. Rediness All']


def _titled(ws, title, width, header_row=3):
    """Row 1: a merged title. Row 2: a merged legend (if the headers are on
    row 3)."""
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    if header_row == 3:
        ws.cell(row=2, column=1, value=' / '.join(TITLE_LEGEND))
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)


def _rows(ws, header_row, headers, rows):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    for r, values in enumerate(rows, start=header_row + 1):
        for c, h in enumerate(headers, start=1):
            if h in values:
                ws.cell(row=r, column=c, value=values[h])


def build(path, extra_sheet=True):
    """Write the made-up workbook to `path`."""
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('Mk2.x')
    _titled(ws, 'PNC Mk2', 12)
    _rows(ws, 3, AIRCRAFT_HEADERS, [
        {'S/N': 'TST-MK2X-0001', 'Configuration': 'Cfg A', 'Frequency': '4.4 (SS5)',
         'Insurance status': '01.01.2027', 'Status': 'RTF', 'Location': 'MRO',
         'Asset owner': 'FOPS', 'INFO UPDATE': datetime(2026, 10, 1), 'SIGN': SIGN,
         'Actions/Defects': '01.10.2026: ' + SECRET_NOTE, 'Equip. Rediness All': '#REF!'},
    ])

    ws = wb.create_sheet('MK2')
    _titled(ws, 'PNC Mk2/2.X', 12)
    headers = [('Owner' if h == 'Asset owner' else h) for h in AIRCRAFT_HEADERS]
    _rows(ws, 3, headers, [
        {'S/N': 'TST-MK2-0001', 'Status': 'GROUNDED', 'Location': 'Field', 'Owner': 'R&D',
         'INFO UPDATE': '15.03.2026', 'SIGN': SIGN},
    ])

    ws = wb.create_sheet('2.5 VTOL')
    _titled(ws, 'PNC Mk2.5 VTOL', 12)
    _rows(ws, 3, AIRCRAFT_HEADERS + ['PRJ202'], [
        {'S/N': 'TST-VTOL-0001', 'Status': 'RTF WITH REMARKS', 'Location': 'Nakotne',
         'Asset owner': 'FOPS', 'INFO UPDATE': 'last spring', 'SIGN': SIGN,
         'Actions/Defects': 'Minor wear on the tail'},
    ])
    # Side table to the right: an upgrade matrix in W to Y.
    for r, values in enumerate([('Upgrade', 'Done', 'Due'), ('Wing', 'x', ''),
                                ('Tail', '', 'x')], start=3):
        for c, v in enumerate(values, start=23):
            ws.cell(row=r, column=c, value=v)
    # Side table below: a part list with nothing in the S/N column.
    for r in range(21, 26):
        ws.cell(row=r, column=2, value='Spare part %d' % r)
        ws.cell(row=r, column=3, value=r - 20)

    ws = wb.create_sheet('MK3')
    _titled(ws, 'PNC Mk3 VTOL', 12)
    _rows(ws, 3, AIRCRAFT_HEADERS, [
        {'S/N': 'TST-MK3-0001', 'Frequency': 4.4, 'Status': 'RTF', 'Location': 'MRO',
         'Asset owner': 'FOPS', 'INFO UPDATE': datetime(2025, 1, 10), 'SIGN': SIGN,
         'Actions/Defects': 'Flown'},
        {'S/N': 'TST-MK3-0002', 'Status': 'FCF REQUIRED', 'Location': '#REF!',
         'Asset owner': 'FOPS', 'INFO UPDATE': datetime(2026, 9, 1), 'SIGN': SIGN,
         'Actions/Defects': 'Awaiting check flight'},
    ])

    ws = wb.create_sheet('VXE')
    _titled(ws, 'PNC VXE', 12)
    headers = ['S/N', 'Configuration', 'Frequency', 'Insurance due', 'STATUS', 'Location',
               'Asset owner', 'INFO UPDATE', 'SIGN', 'Actions/Defects', 'NOTES']
    _rows(ws, 3, headers, [
        {'S/N': 4417, 'Configuration': 'VXE 30 - Battery cell', 'STATUS': 'RTF',
         'Location': 'DEMO', 'Asset owner': 'FOPS', 'INFO UPDATE': '20.10.2026', 'SIGN': SIGN,
         'NOTES': 'Demo unit'},
    ])

    ws = wb.create_sheet('EPS')
    _titled(ws, 'EPSILON GIMBALS', 10)
    headers = ['S/N', 'Configuration', 'Positioned in', 'Location', 'Asset Owner', 'STATUS',
               'SIGN', 'Date', 'Actions/Defects', 'NC #']
    _rows(ws, 3, headers, [
        {'S/N': 'TST-EPS-0001', 'Configuration': 'E180', 'Positioned in': 'TST-MK3-0001',
         'Location': 'MRO', 'Asset Owner': 'FOPS', 'STATUS': 'NFC ISSUES', 'SIGN': SIGN,
         'Date': '02.09.2026', 'Actions/Defects': 'Gimbal drift'},
    ])

    ws = wb.create_sheet('TRC')
    _titled(ws, 'TRACKING ANTENNA', 11)
    headers = ['S/N', 'Configuration', 'Frequency', 'Status', 'Location', 'Asset Owner', 'SIGN',
               'Licences', 'NC #', 'DATE', 'Actions/Defects']
    _rows(ws, 3, headers, [
        {'S/N': 'TST-TRC-0001', 'Frequency': 2.4, 'Status': 'RTF', 'Location': 'Field',
         'Asset Owner': 'FOPS', 'SIGN': SIGN, 'DATE': '05.10.2026'},
        {'S/N': 'TST-TRC-0002', 'Frequency': 2.4, 'Status': 'RTF', 'Location': 'Field',
         'Asset Owner': 'FOPS', 'SIGN': SIGN, 'DATE': '05.10.26'},
        {'S/N': 'TST-TRC-0003', 'Status': 'GROUNDED', 'Location': 'MRO',
         'Asset Owner': 'R&D', 'SIGN': SIGN, 'Actions/Defects': 'Cable cut'},
    ])
    # Riga's own count table beside the data, in M and N.
    for r, (label, n) in enumerate([('RTF', 2), ('GROUNDED', 1), ('GRTF WITH REMARKS', 0)],
                                   start=3):
        ws.cell(row=r, column=13, value=label)
        ws.cell(row=r, column=14, value=n)

    ws = wb.create_sheet('PPC')
    _titled(ws, 'Catapult', 9, header_row=2)
    headers = ['S/N', 'Configuration', 'Location', 'Asset Owner', 'Status', 'SIGN',
               'NEXT MNT (launches)', 'Date', 'INFO']
    _rows(ws, 2, headers, [
        {'S/N': 'TST-PPC-0001', 'Configuration': '4m', 'Location': 'Field', 'Asset Owner': 'FOPS',
         'Status': 'RTF', 'SIGN': SIGN, 'NEXT MNT (launches)': 40, 'Date': datetime(2026, 10, 20),
         'INFO': 'New bungee'},
    ])

    ws = wb.create_sheet('GCS')
    _titled(ws, 'GROUND CONTROL STATIONS', 10)
    headers = ['S/N', 'STATUS', 'Configuration', 'PC INSTALLED', 'Location', 'Asset Owner',
               'INFO:', 'SIGN', 'Actions/Defects', 'NC #']
    _rows(ws, 3, headers, [
        {'S/N': 'TST-GCS-0001', 'STATUS': 'RTF', 'Configuration': 'std', 'Location': 'MRO',
         'Asset Owner': 'FOPS'},
    ])

    ws = wb.create_sheet('ENG')
    _titled(ws, 'ENGINES', 11)
    headers = ['S/N', 'Configuration', 'Positioned in', 'Asset Owner', 'Status', 'Location',
               'SIGN', 'INFO UPDATE', 'Actions/Defects', 'Planned RTF', 'NC #']
    _rows(ws, 3, headers, [
        {'S/N': 'TST-ENG-0001', 'Status': 'RTF', 'Location': 'MRO', 'Asset Owner': 'FOPS',
         'INFO UPDATE': datetime(2026, 10, 2), 'SIGN': SIGN, 'Positioned in': 'On shelf'},
        {'S/N': 'TST-ENG-0002', 'Status': 'RTF', 'Location': 'Field', 'Asset Owner': 'FOPS',
         'INFO UPDATE': '03.10.2026', 'SIGN': SIGN},
        {'S/N': 'TST-ENG-0003', 'Status': 'GROUNDED', 'Location': 'MRO', 'Asset Owner': 'FOPS',
         'INFO UPDATE': datetime(2026, 10, 4), 'SIGN': SIGN, 'Actions/Defects': 'Low compression'},
        {'S/N': 'TST-ENG-0004', 'Status': 'GROUNDED', 'Location': 'MRO', 'Asset Owner': 'R&D',
         'SIGN': SIGN},
        {'S/N': 'TST-ENG-0005', 'Status': 'WRITTEN OFF', 'Location': 'MRO', 'Asset Owner': 'FOPS',
         'INFO UPDATE': datetime(2026, 6, 1), 'SIGN': SIGN},
    ])
    # The unlabelled marker columns L and M.
    for r in range(4, 9):
        ws.cell(row=r, column=12, value='X')
    ws.cell(row=4, column=13, value='RTF')

    ws = wb.create_sheet('VEHICLES')
    _titled(ws, 'FLOPS CAR UNIT', 4, header_row=2)
    _rows(ws, 2, ['Car', 'Number plate', 'KASKO'], [
        {'Car': 'Test car', 'Number plate': 'XX-0000', 'KASKO': '01.02.2027'}])

    ws = wb.create_sheet('ZALENIEKI STOCK')
    _titled(ws, 'Critical Spare Part List', 4, header_row=2)
    _rows(ws, 2, ['Name', 'Required Quantity', 'P/N', 'Available QTY'], [
        {'Name': 'Test prop', 'Required Quantity': 4, 'P/N': 'PN-TEST', 'Available QTY': 2}])

    ws = wb.create_sheet('OVERALL')
    _titled(ws, 'STATUS TABLE (READ ONLY)', 4, header_row=2)
    _rows(ws, 2, ['Model', 'S/N', 'Status'], [
        {'Model': 'Mk3', 'S/N': '#REF!', 'Status': '#REF!'}])

    if extra_sheet:
        ws = wb.create_sheet('Notes 2026')
        ws.cell(row=1, column=1, value='Meeting notes')

    wb.save(path)
    return path


# What the fixture holds, for the tests.
ROWS_PER_SHEET = {'Mk2.x': 1, 'MK2': 1, '2.5 VTOL': 1, 'MK3': 2, 'VXE': 1, 'EPS': 1,
                  'TRC': 3, 'PPC': 1, 'GCS': 1, 'ENG': 5}
SERIALS = ['TST-MK2X-0001', 'TST-MK2-0001', 'TST-VTOL-0001', 'TST-MK3-0001', 'TST-MK3-0002',
           '4417', 'TST-EPS-0001', 'TST-TRC-0001', 'TST-TRC-0002', 'TST-TRC-0003',
           'TST-PPC-0001', 'TST-GCS-0001', 'TST-ENG-0001', 'TST-ENG-0002', 'TST-ENG-0003',
           'TST-ENG-0004', 'TST-ENG-0005']
