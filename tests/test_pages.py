import io
import logging
import re
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

import app as dashboard
import sharepoint
from fake_sharepoint import FakeClock, FakeSharePoint, FakeTokens, fake_env, make_item

PAGES = ['/', '/inventory', '/attention', '/checkouts', '/metrics', '/api/summary', '/api/assets',
         '/api/attention', '/api/fields', '/api/export/assets.xlsx']
PEOPLE = {'7': {'Title': 'Pat Example', 'EMail': 'pat.example@example.com'}}


def items():
    return [
        make_item(17, 'SN-FAKE-017', status=['OUT'], person=7, date_out='2026-09-03T07:00:00Z'),
        make_item(23, 'SN-FAKE-023', condition='NMC', last_inventory='2026-01-27T08:00:00Z'),
        make_item(42, 'SN-FAKE-042', status=['RMA'], location=['Test Bay B']),
        # Checked out to someone but not marked OUT, and OUT with nobody listed.
        make_item(51, 'SN-FAKE-051', status=['IN'], person=7),
        make_item(52, 'SN-FAKE-052', status=['OUT'], date_out='2026-09-03T07:00:00Z'),
    ]


@pytest.fixture
def sharepoint_app(pfx, monkeypatch):
    fake = FakeSharePoint(items(), people=PEOPLE)
    clock = FakeClock()
    settings = sharepoint.load_settings(fake_env(pfx))
    client = sharepoint.GraphClient(FakeTokens(), session=fake, sleep=lambda s: None)
    source = sharepoint.SharePointSource(settings, 'America/Los_Angeles', client=client,
                                         clock=clock, background=False)
    monkeypatch.setattr(dashboard, 'SOURCE', source)
    return dashboard.app.test_client(), fake, clock


@pytest.mark.parametrize('path', PAGES)
def test_every_page_with_the_csv_source(path):
    assert dashboard.app.test_client().get(path).status_code == 200


@pytest.mark.parametrize('path', PAGES)
def test_every_page_with_the_sharepoint_source(sharepoint_app, path):
    client, _, _ = sharepoint_app
    assert client.get(path).status_code == 200


def test_pages_show_sharepoint_data_and_item_ids(sharepoint_app):
    client, _, _ = sharepoint_app
    html = client.get('/inventory').get_data(as_text=True)
    assert 'SharePoint: Asset Tracking V2' in html and ', fetched ' in html
    for item_id in ('17', '23', '42'):
        assert 'data-id="%s"' % item_id in html
    ids = [a['id'] for a in client.get('/api/assets').get_json()['assets']]
    assert ids == [17, 23, 42, 51, 52]


def test_names_on_every_view_with_the_sharepoint_source(sharepoint_app):
    client, _, _ = sharepoint_app
    checkouts = client.get('/checkouts').get_data(as_text=True)
    assert 'Pat Example' in checkouts and 'pat.example@example.com' not in checkouts
    assert 'Pat Example' in client.get('/inventory').get_data(as_text=True)
    # The Attention rules that test Check Out By/To work the same with names.
    groups = {g['key']: [a['id'] for a in g['assets']]
              for g in client.get('/api/attention').get_json()}
    assert groups['holder_not_out'] == [51]
    assert groups['out_no_holder'] == [52]
    holders = [a['checked_out_to'] for a in client.get('/api/assets').get_json()['assets']]
    assert holders == ['Pat Example', '', '', 'Pat Example', '']
    ws = load_workbook(io.BytesIO(client.get('/api/export/assets.xlsx?ids=17').data)).active
    assert ws.cell(row=1, column=9).value == 'Check Out By/To'
    assert ws.cell(row=2, column=9).value == 'Pat Example'


def test_csv_source_still_shows_what_the_export_holds():
    rows = dashboard.app.test_client().get('/api/assets').get_json()['assets']
    holders = {a['checked_out_to'] for a in rows if a['checked_out_to']}
    assert holders and all('@' in h for h in holders)


def test_excel_export_ids_filter_uses_sharepoint_item_ids(sharepoint_app):
    client, _, _ = sharepoint_app
    resp = client.get('/api/export/assets.xlsx?ids=17,42')
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    serials = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert serials == ['SN-FAKE-017', 'SN-FAKE-042']


def test_stale_copy_shows_a_short_note_and_the_reason_only_in_the_log(sharepoint_app, caplog):
    client, fake, clock = sharepoint_app
    assert client.get('/inventory').status_code == 200
    fake.fail_status = 503
    clock.advance(301)
    with caplog.at_level(logging.INFO):
        client.get('/inventory')
    since = dashboard.SOURCE.load()['stale']['since']
    expected = ('Data is from %s. SharePoint could not be reached; showing the last copy.'
                % dashboard._clock(since))
    for path in ['/', '/inventory', '/attention', '/checkouts', '/metrics']:
        resp = client.get(path)
        html = resp.get_data(as_text=True)
        assert resp.status_code == 200
        assert expected in html
        assert 'server problem' not in html and 'HTTP 503' not in html
    assert 'SN-FAKE-017' in client.get('/inventory').get_data(as_text=True)
    assert 'server problem' in caplog.text and 'HTTP 503' in caplog.text


def test_stale_note_gives_the_date_when_the_copy_is_from_an_earlier_day():
    yesterday = datetime.now(ZoneInfo('America/Los_Angeles')) - timedelta(days=1)
    yesterday = yesterday.replace(hour=9, minute=40)
    assert dashboard.short_time(yesterday) == '%d/%d/%d 9:40 AM' % (
        yesterday.month, yesterday.day, yesterday.year)
    today_ = yesterday + timedelta(days=1)
    assert dashboard.short_time(today_.replace(hour=14, minute=5)) == '2:05 PM'


def test_no_good_fetch_says_sharepoint_could_not_be_reached(pfx, monkeypatch, caplog):
    fake = FakeSharePoint(items())
    fake.fail_status = 403
    settings = sharepoint.load_settings(fake_env(pfx))
    client = sharepoint.GraphClient(FakeTokens(), session=fake, sleep=lambda s: None)
    source = sharepoint.SharePointSource(settings, 'America/Los_Angeles', client=client,
                                         clock=FakeClock(), background=False)
    monkeypatch.setattr(dashboard, 'SOURCE', source)
    with caplog.at_level(logging.INFO):
        resp = dashboard.app.test_client().get('/inventory')
    html = resp.get_data(as_text=True)
    assert resp.status_code == 503
    assert 'SharePoint could not be reached' in html
    assert 'Permission missing' not in html and 'HTTP 403' not in html
    assert 'Permission missing' in caplog.text
    api = dashboard.app.test_client().get('/api/assets')
    assert api.status_code == 503
    assert 'SharePoint could not be reached' in api.get_json()['error']


# --- A bad SharePoint setting keeps the dashboard up -------------------------

BAD_SETTINGS = {
    'missing setting': ({'SP_SITE_ID': None}, 'SP_SITE_ID'),
    'missing certificate file': ({'SP_CERT_PATH': 'C:\\nowhere\\secret-name.pfx'}, 'SP_CERT_PATH'),
    'wrong password': ({'SP_CERT_PASSWORD': 'hunter2-wrong'}, 'SP_CERT_PASSWORD'),
    'bad refresh': ({'SP_REFRESH_SECONDS': '5'}, 'SP_REFRESH_SECONDS'),
}


@pytest.mark.parametrize('case', sorted(BAD_SETTINGS))
def test_bad_setting_keeps_the_app_up_with_an_error_on_every_page(pfx, monkeypatch, caplog, case):
    overrides, setting = BAD_SETTINGS[case]
    env = fake_env(pfx, **overrides)
    with caplog.at_level(logging.INFO):
        source = dashboard.make_source('sharepoint', env)
        monkeypatch.setattr(dashboard, 'SOURCE', source)
        client = dashboard.app.test_client()
        bodies = []
        for path in PAGES:
            resp = client.get(path)
            assert resp.status_code == 503, path
            if path.startswith('/api/') and 'export' not in path:
                bodies.append(resp.get_json()['error'])
            else:
                html = resp.get_data(as_text=True)
                assert 'The asset list cannot be shown right now' in html
                bodies.append(html)
    for body in bodies:
        assert setting in body
        assert 'settings problem' in body
        # No setting's value is printed (short ones like "5" would match any page).
        for name, value in env.items():
            if name.startswith('SP_') and name != 'SP_LIST_NAME' and len(value) > 3:
                assert value not in body, name
        assert 'hunter2-wrong' not in body and 'secret-name' not in body
    # Logged once, at startup, not on each of the requests.
    problems = [r for r in caplog.records if 'Settings problem' in r.getMessage()]
    assert len(problems) == 1 and setting in problems[0].getMessage()
    assert 'hunter2-wrong' not in caplog.text and pfx['password'] not in caplog.text


def test_unknown_asset_source_keeps_the_app_up(monkeypatch):
    source = dashboard.make_source('sharepiont')
    monkeypatch.setattr(dashboard, 'SOURCE', source)
    resp = dashboard.app.test_client().get('/inventory')
    assert resp.status_code == 503
    assert 'ASSET_SOURCE must be csv or sharepoint' in resp.get_data(as_text=True)


# --- Addresses and the left nav ---------------------------------------------

HTML_PAGES = {'/': 'overview', '/inventory': 'inventory', '/attention': 'attention',
              '/checkouts': 'checkouts', '/metrics': 'metrics'}


def test_home_page_is_the_overview_and_inventory_moved():
    client = dashboard.app.test_client()
    assert 'Overall Fleet Readiness' in client.get('/').get_data(as_text=True)
    inventory = client.get('/inventory').get_data(as_text=True)
    assert 'id="asset-table"' in inventory and 'Overall Fleet Readiness' not in inventory


@pytest.mark.parametrize('old, new', [
    ('/overview', '/'),
    ('/overview?location=SLO+ENG', '/?location=SLO+ENG'),
    ('/overview?location=%28none%29', '/?location=%28none%29'),
])
def test_old_overview_address_redirects_with_302(old, new):
    resp = dashboard.app.test_client().get(old)
    assert resp.status_code == 302
    assert resp.headers['Location'] == new


def links(html):
    return re.findall(r'href="([^"]*)"', html)


@pytest.mark.parametrize('path', sorted(HTML_PAGES))
def test_no_page_links_to_the_home_page_meaning_inventory(path):
    html = dashboard.app.test_client().get(path).get_data(as_text=True)
    assert '/overview' not in links(html)
    # The only links to "/" are the nav's Fleet Overview and the overview's own
    # "Show all"; the Metrics Assets card goes to /inventory.
    for m in re.finditer(r'<a [^>]*href="/"[^>]*>([^<]*)<', html):
        assert m.group(1).strip() in ('Fleet Overview', 'Show all'), m.group(0)
    if path == '/metrics':
        assert '<a class="kpi-card" href="/inventory">' in html


@pytest.mark.parametrize('path', sorted(HTML_PAGES))
def test_every_page_has_the_left_nav_with_its_page_highlighted(path):
    html = dashboard.app.test_client().get(path).get_data(as_text=True)
    assert 'class="side-nav"' in html and 'header-bar' not in html
    assert 'id="theme-toggle"' in html
    nav = html.split('<nav class="side-links">', 1)[1].split('</nav>', 1)[0]
    assert re.findall(r'href="([^"]+)"', nav) == [href for _, href, _ in dashboard.PAGES]
    active = re.findall(r'<a href="([^"]+)" class="active" aria-current="page">', nav)
    assert active == [href for key, href, _ in dashboard.PAGES if key == HTML_PAGES[path]]


def test_excel_export_button_is_on_the_inventory_page():
    html = dashboard.app.test_client().get('/inventory').get_data(as_text=True)
    head = html.split('<div class="page-head">', 1)[1].split('</div>\n', 2)
    assert 'id="export-link" href="/api/export/assets.xlsx">Export to Excel</a>' in ''.join(head[:2])


def test_error_page_has_the_left_nav(monkeypatch):
    monkeypatch.setattr(dashboard, 'SOURCE', dashboard.UnavailableSource('Broken for the test.'))
    html = dashboard.app.test_client().get('/').get_data(as_text=True)
    assert 'class="side-nav"' in html and 'class="active"' not in html
