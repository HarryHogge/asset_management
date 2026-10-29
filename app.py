from flask import Flask, render_template, jsonify, request, send_file, abort, redirect
from datetime import datetime
from dotenv import load_dotenv
from zoneinfo import ZoneInfo
import io
import logging
import os

import assets as asset_data
import fleet
import riga_loader
import sharepoint

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, '.env'), encoding='utf-8')

if __name__ == '__main__':
    # Set up before the data source is made, so a settings problem found at
    # startup reaches the log in the same format as everything else. MSAL is
    # held to warnings so the log never fills with sign-in chatter.
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(name)s: %(message)s')
    logging.getLogger('msal').setLevel(logging.WARNING)
log = logging.getLogger('asset_management.app')

app = Flask(__name__)

# Where the SharePoint export lives. Relative paths resolve against this repo.
ASSET_CSV = os.getenv('ASSET_CSV', 'Asset Tracking V2.csv')
if not os.path.isabs(ASSET_CSV):
    ASSET_CSV = os.path.join(BASE_DIR, ASSET_CSV)
# The SharePoint site's time zone; list dates are stored as local midnight in UTC.
SITE_TIMEZONE = os.getenv('SITE_TIMEZONE', 'America/Los_Angeles')
INVENTORY_STALE_DAYS = int(os.getenv('INVENTORY_STALE_DAYS', '90'))
# Not 5060/5061: those are SIP ports and Chrome refuses to open them.
PORT = int(os.getenv('PORT', '5070'))


def history_path(env=None):
    """READINESS_HISTORY_PATH: the file holding one line per day of readiness
    numbers, for the Fleet Overview trend. Relative paths resolve against
    this repo."""
    env = os.environ if env is None else env
    path = env.get('READINESS_HISTORY_PATH') or os.path.join('data', 'readiness_history.jsonl')
    return path if os.path.isabs(path) else os.path.join(BASE_DIR, path)


def make_history(env=None):
    """The readiness history, with its folder made now. If the folder cannot
    be made, history is off until the dashboard restarts: logged once here,
    and /overview shows a short note."""
    env = os.environ if env is None else env
    if 'READINESS_HISTORY' in env:
        log.warning('The setting READINESS_HISTORY is ignored: it was renamed '
                    'READINESS_HISTORY_PATH. Rename it in .env.')
    path = history_path(env)
    folder = os.path.dirname(path)
    try:
        os.makedirs(folder, exist_ok=True)
    except OSError as e:
        reason = 'the folder %s could not be created (%s: %s)' % (folder, type(e).__name__, e)
        log.error('Readiness history is OFF until the dashboard is restarted, because %s. '
                  'Fix READINESS_HISTORY_PATH in .env, then restart.', reason)
        return fleet.ReadinessHistory(path, off_reason=reason)
    return fleet.ReadinessHistory(path)


def riga_workbook_path(env=None):
    """RIGA_WORKBOOK_PATH: where the uploaded Riga workbook is kept. Relative
    paths resolve against this repo."""
    env = os.environ if env is None else env
    path = env.get('RIGA_WORKBOOK_PATH') or os.path.join('data', 'riga', 'current.xlsx')
    return path if os.path.isabs(path) else os.path.join(BASE_DIR, path)


def make_riga(env=None):
    """The Riga workbook source, with its folder made now, as for the history.
    If the folder cannot be made, the Riga view and the upload page say so
    until a restart, and the US view is not affected."""
    path = riga_workbook_path(env)
    folder = os.path.dirname(path)
    try:
        os.makedirs(folder, exist_ok=True)
    except OSError as e:
        reason = 'the folder %s could not be created (%s: %s)' % (folder, type(e).__name__, e)
        log.error('The Riga workbook is OFF until the dashboard is restarted, because %s. '
                  'Fix RIGA_WORKBOOK_PATH in .env, then restart.', reason)
        return riga_loader.RigaSource(path, off_reason=reason)
    return riga_loader.RigaSource(path)


class CsvSource:
    """The SharePoint list's CSV export (the default source)."""

    def __init__(self, path, tz_name):
        self.path = path
        self.tz_name = tz_name

    def load(self):
        return asset_data.load(self.path, self.tz_name)


class UnavailableSource:
    """Stands in for the data source when a setting is wrong. The dashboard
    stays up (under PM2, exiting would only restart it over and over) and
    every page says which setting to fix."""

    def __init__(self, problem):
        self.problem = problem

    def load(self):
        raise sharepoint.ConfigError(self.problem)


def make_source(name, env=None):
    """The data source named by ASSET_SOURCE. A settings problem is logged
    here, once, and every page then shows it until .env is fixed."""
    name = (name or 'csv').strip().lower()
    try:
        if name == 'csv':
            return CsvSource(ASSET_CSV, SITE_TIMEZONE)
        if name == 'sharepoint':
            return sharepoint.SharePointSource.from_env(SITE_TIMEZONE, env)
        raise sharepoint.ConfigError('ASSET_SOURCE must be csv or sharepoint.')
    except sharepoint.ConfigError as e:
        log.error('Settings problem: %s The dashboard is running, but every page shows '
                  'this until .env is fixed and the dashboard is restarted.', e)
        return UnavailableSource(str(e))


SOURCE = make_source(os.getenv('ASSET_SOURCE'))
FAMILIES = fleet.FamilyMap()
HISTORY = make_history()
RIGA = make_riga()

# Requests larger than this are cut off before they reach a page. The Riga
# upload checks its own, smaller limit and says so on the page; this is the
# outer guard, with room for the rest of the form.
app.config['MAX_CONTENT_LENGTH'] = riga_loader.MAX_UPLOAD_BYTES + 1024 * 1024

PAGES = [
    ('overview', '/', 'Fleet Overview'),
    ('inventory', '/inventory', 'Inventory'),
    ('attention', '/attention', 'Needs Attention'),
    ('checkouts', '/checkouts', 'Checkouts'),
    ('metrics', '/metrics', 'Metrics'),
]


def today():
    return datetime.now(ZoneInfo(SITE_TIMEZONE)).date()


def riga_now():
    return datetime.now(ZoneInfo(riga_loader.RIGA_TIMEZONE))


# The family cross-check problems last logged, so each is logged once.
_riga_logged = {'mismatches': ()}


def load_riga():
    """The Riga workbook as RigaSource.load gives it, never raising. The
    first good load of each Riga day adds Riga's line to the history."""
    try:
        state = RIGA.load()
    except Exception as e:  # a bug here must never break the US view
        log.exception('Riga workbook: unexpected error while loading.')
        return {'state': 'unreadable', 'problem': '%s: %s' % (type(e).__name__, e)}
    if state['state'] == 'ok':
        mapping, problem = FAMILIES.load()
        use, mismatches = riga_loader.cross_check(mapping, state['mapping'])
        state['families'] = use
        # With the family file itself broken, its own note covers it.
        state['mismatches'] = [] if problem else mismatches
        key = tuple(state['mismatches'])
        if key != _riga_logged['mismatches']:
            _riga_logged['mismatches'] = key
            for m in key:
                log.warning('Riga workbook: %s', riga_loader.mismatch_text(m))
        HISTORY.record(state['data'], None, riga_now(), site=fleet.RIGA, mapping=use)
    return state


def load():
    try:
        data = SOURCE.load()
    except FileNotFoundError:
        abort(503, description='Asset export not found: %s (set ASSET_CSV in .env)'
              % os.path.basename(ASSET_CSV))
    except sharepoint.ConfigError as e:
        abort(503, description='The dashboard cannot read the asset list because of a '
              'settings problem. %s Fix it in .env on the server, then restart the '
              'dashboard.' % e)
    except sharepoint.SourceError:
        # The reason is already in the log, written when the fetch failed.
        abort(503, description='SharePoint could not be reached, and there is no earlier '
              'copy of the asset list to show. The dashboard tries again every minute, '
              'so reload this page shortly. The reason is in the dashboard log.')
    # The first good load of each day adds that day's line to the readiness
    # history. It never raises: a failed write is logged and noted on /overview.
    HISTORY.record(data, FAMILIES, datetime.now(ZoneInfo(SITE_TIMEZONE)))
    return data


@app.errorhandler(503)
def unavailable(e):
    # The JSON endpoints answer in JSON; pages and the Excel download get a page.
    if request.path.startswith('/api/') and not request.path.startswith('/api/export/'):
        return jsonify({'error': e.description}), 503
    return render_template('error.html', active=None, message=e.description), 503


@app.context_processor
def page_context():
    return {'pages': PAGES, 'stale_days': INVENTORY_STALE_DAYS}


# A global, not context: _macros.html is imported without the page's context.
app.jinja_env.globals['light_pill'] = asset_data.light_pill
app.jinja_env.globals['mismatch_text'] = riga_loader.mismatch_text


@app.template_filter('mdy')
def mdy(value):
    return '%d/%d/%d' % (value.month, value.day, value.year) if value else ''


def _clock(value):
    # Built by hand rather than with strftime's %-m/%-I: those no-leading-zero
    # flags don't exist on Windows and raise ValueError there.
    hour = value.hour % 12 or 12
    return '%d:%02d %s' % (hour, value.minute, 'AM' if value.hour < 12 else 'PM')


@app.template_filter('mdy_time')
def mdy_time(value):
    return '%s %s' % (mdy(value), _clock(value)) if value else ''


@app.template_filter('short_time')
def short_time(value):
    """Just the time if it is today in the site's time zone, otherwise the
    date and time, so "9:40 AM" never means yesterday."""
    if not value:
        return ''
    if value.astimezone(ZoneInfo(SITE_TIMEZONE)).date() == today():
        return _clock(value)
    return mdy_time(value)


def render(template, active, data, **kwargs):
    # `data` is the copy the route already used, so a background refresh that
    # finishes mid-request cannot mix two versions on one page.
    return render_template(
        template,
        active=active,
        fields=data['fields'],
        source=data['source'],
        loaded_at=data['loaded_at'],
        live=data.get('live', False),
        stale=data.get('stale'),
        **kwargs,
    )


@app.route('/')
def overview():
    # The Riga workbook is loaded for either site's view, so Riga's trend
    # fills in even on days nobody opens the Riga view.
    riga = load_riga()
    if request.args.get('site', '').strip().lower() == 'riga':
        return riga_overview(riga)
    data = load()
    mapping, mapping_problem = FAMILIES.load()
    location = request.args.get('location', '')
    assets = fleet.filter_location(data['assets'], location)
    locations, any_without = fleet.location_options(data['assets'], data['fields'])
    groups = asset_data.attention(assets, today(), INVENTORY_STALE_DAYS)
    alert_rows, problem_count = fleet.alert_rows(groups)
    # The trend is site-wide: it comes from the history, not the filter.
    trend = fleet.trend(HISTORY.entries(fleet.US), today())
    return render('overview.html', 'overview', data,
                  site='us',
                  show_checkouts=True,
                  location=location,
                  no_location=fleet.NO_LOCATION,
                  locations=locations,
                  any_without_location=any_without,
                  ready=fleet.fleet_readiness(assets, mapping),
                  tables=[(f, fleet.table_rows(assets, mapping, f)) for f in fleet.MAIN_FAMILIES],
                  unmapped=fleet.unmapped_types(assets, mapping),
                  mapping_problem=mapping_problem,
                  history_off=bool(HISTORY.off_reason),
                  history_problem=HISTORY.problem,
                  alert_rows=alert_rows,
                  problem_count=problem_count,
                  trend=trend)


def riga_overview(riga):
    """The overview for Riga: the same panels, from the uploaded workbook.
    Riga has no checkout data, so Status and Checked out to are left out,
    and engines get a table of their own."""
    location = request.args.get('location', '')
    common = dict(site='riga', show_checkouts=False, location=location,
                  no_location=fleet.NO_LOCATION, riga=riga,
                  history_off=bool(HISTORY.off_reason), history_problem=HISTORY.problem)
    if riga['state'] != 'ok':
        return render_template('overview.html', active='overview', fields={}, source=None,
                               **common)
    data = riga['data']
    mapping = riga['families']
    _, mapping_problem = FAMILIES.load()
    riga_today = riga_now().date()
    assets = fleet.filter_location(data['assets'], location)
    locations, any_without = fleet.location_options(data['assets'], data['fields'])
    groups = riga_loader.attention(assets, riga_today, INVENTORY_STALE_DAYS,
                                   riga['result']['dated_types'])
    alert_rows, problem_count = fleet.alert_rows(groups)
    trend = fleet.trend(HISTORY.entries(fleet.RIGA), riga_today)
    engines = fleet.table_rows(assets, mapping, fleet.ENGINE)
    return render('overview.html', 'overview', data,
                  locations=locations,
                  any_without_location=any_without,
                  ready=fleet.fleet_readiness(assets, mapping),
                  tables=[(f, fleet.table_rows(assets, mapping, f)) for f in fleet.MAIN_FAMILIES],
                  engines=engines,
                  engine_ready=fleet.readiness(engines),
                  unmapped=fleet.unmapped_types(assets, mapping),
                  mapping_problem=mapping_problem,
                  alert_rows=alert_rows,
                  problem_count=problem_count,
                  trend=trend,
                  **common)


@app.route('/riga/upload', methods=['GET', 'POST'])
def riga_upload():
    """Upload a new Riga workbook. Anyone who can open the dashboard can use
    this page; the name is whatever they type."""
    message, error, status = None, None, 200
    if request.method == 'POST':
        f = request.files.get('workbook')
        try:
            info = riga_loader.save_upload(
                RIGA, f.stream if f else None, f.filename if f else '',
                request.form.get('name', ''), riga_now())
            message = info
        except riga_loader.UploadRejected as e:
            error, status = str(e), 400
    return riga_upload_page(message=message, error=error), status


def riga_upload_page(message=None, error=None):
    return render_template('riga_upload.html', active='overview', source=None,
                           off_reason=RIGA.off_reason, upload=RIGA.upload_info(),
                           limit=riga_loader.size_text(riga_loader.MAX_UPLOAD_BYTES),
                           max_name=riga_loader.MAX_NAME_LENGTH,
                           message=message, error=error)


@app.errorhandler(413)
def too_large(e):
    if request.path == '/riga/upload':
        return riga_upload_page(error='The file is larger than the %s limit.'
                                % riga_loader.size_text(riga_loader.MAX_UPLOAD_BYTES)), 413
    return render_template('error.html', active=None, message='The request was too large.'), 413


@app.route('/overview')
def overview_moved():
    """The overview's old address. A temporary redirect, so browsers do not
    remember it if the home page changes again; the query string is kept."""
    query = request.query_string.decode('utf-8', 'replace')
    return redirect('/' + ('?' + query if query else ''), code=302)


@app.route('/inventory')
def inventory():
    data = load()
    groups = asset_data.attention(data['assets'], today(), INVENTORY_STALE_DAYS)
    return render('inventory.html', 'inventory', data,
                  assets=data['assets'],
                  options=asset_data.filter_options(data['assets'], data['fields']),
                  flagged=asset_data.flagged_ids(groups))


@app.route('/attention')
def attention():
    data = load()
    groups = asset_data.attention(data['assets'], today(), INVENTORY_STALE_DAYS)
    return render('attention.html', 'attention', data, groups=groups)


@app.route('/checkouts')
def checkouts():
    data = load()
    rows = asset_data.checkouts(data['assets'], today())
    holders = {}
    for r in rows:
        holders.setdefault(r['checked_out_to'] or '(nobody listed)', []).append(r)
    holders = sorted(holders.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    return render('checkouts.html', 'checkouts', data, rows=rows, holders=holders)


@app.route('/metrics')
def metrics():
    data = load()
    return render('metrics.html', 'metrics', data,
                  summary=asset_data.summary(data['assets'], today(), INVENTORY_STALE_DAYS))


@app.route('/api/assets')
def api_assets():
    data = load()
    return jsonify({
        'source': data['source'],
        'loaded_at': data['loaded_at'].isoformat(),
        'count': len(data['assets']),
        'assets': [asset_data.serialize(a) for a in data['assets']],
    })


@app.route('/api/fields')
def api_fields():
    return jsonify(load()['fields'])


@app.route('/api/summary')
def api_summary():
    return jsonify(asset_data.summary(load()['assets'], today(), INVENTORY_STALE_DAYS))


@app.route('/api/attention')
def api_attention():
    groups = asset_data.attention(load()['assets'], today(), INVENTORY_STALE_DAYS)
    return jsonify([
        dict(g, assets=[asset_data.serialize(a) for a in g['assets']]) for g in groups
    ])


@app.route('/api/export/assets.xlsx')
def export_assets():
    """Excel export of the inventory, optionally limited to ?ids=1,2,3
    (the inventory page passes whatever its filters currently show)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    data = load()
    rows = data['assets']
    ids = request.args.get('ids')
    if ids:
        wanted = {int(i) for i in ids.split(',') if i.strip().isdigit()}
        rows = [a for a in rows if a['id'] in wanted]

    columns = [
        ('Asset Type', 'asset_type', 20), ('Serial Number', 'serial_number', 16),
        ('Part Number', 'part_number', 18), ('Ownership', 'ownership', 20),
        ('Status', 'status', 10), ('Location', 'location', 22),
        ('Material Condition', 'condition', 12), ('Date Out', 'date_out', 12),
        ('Check Out By/To', 'checked_out_to', 28), ('Shelf Location', 'shelf_location', 24),
        ('Last Inventory', 'last_inventory', 14), ('Asset Notes', 'notes', 50),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Assets'
    header_fill = PatternFill('solid', fgColor='B00000')
    for col, (label, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, asset in enumerate(rows, start=2):
        for col, (_, key, _) in enumerate(columns, start=1):
            value = asset[key]
            if isinstance(value, list):
                value = ', '.join(value)
            cell = ws.cell(row=r, column=col, value=value)
            if key in asset_data.DATES and value:
                cell.number_format = 'm/d/yyyy'
            if key == 'notes':
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='assets-%s.xlsx' % today().isoformat(),
    )


if __name__ == '__main__':
    from waitress import serve
    print("=" * 60)
    print("Asset Tracking Dashboard starting...")
    if isinstance(SOURCE, CsvSource):
        print(f"Data source: CSV export {ASSET_CSV}")
    elif isinstance(SOURCE, UnavailableSource):
        print("Data source: NOT AVAILABLE because of a settings problem (see the log line")
        print("above). Every page shows the problem until .env is fixed and the dashboard")
        print("is restarted.")
    else:
        print(f"Data source: SharePoint list {SOURCE.settings['list_name']}, "
              f"refreshed every {SOURCE.settings['refresh_seconds']} seconds")
    print(f"Site time zone: {SITE_TIMEZONE}")
    if HISTORY.off_reason:
        print(f"Readiness history: OFF until restart, because {HISTORY.off_reason}")
    else:
        print(f"Readiness history: {HISTORY.path}")
    if RIGA.off_reason:
        print(f"Riga workbook: OFF until restart, because {RIGA.off_reason}")
    else:
        print(f"Riga workbook: {RIGA.path}")
    print("=" * 60)
    print(f"\nOpen http://localhost:{PORT} in your browser\n")
    serve(app, host='0.0.0.0', port=PORT)
