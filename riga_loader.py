"""Read the Riga site's fleet workbook (an Excel file with one sheet per kind
of equipment) into the same rows the US list gives.

Every Riga-specific rule is in riga_mapping.json, not here: which sheets are
read and which are ignored, each sheet's header row and header names, what
each Riga Status means, and which text date formats are accepted. If the
real workbook differs from the sample report 10 describes, the mapping file
is edited and this code stays as it is.

Only the mapped columns of each sheet's header row are read, from the row
below it, and only where the serial number is filled in. Titles, legends and
side tables are skipped because of where they are, not because their content
is recognised. Cached values are read, never formulas.

Check any workbook, with no dashboard running and no SharePoint settings:

    python -m riga_loader --check PATH

Its output holds no full serial numbers, notes, initials or SIGN values, so
it is safe to paste into a chat.
"""
import json
import logging
import os
import re
import shutil
import sys
import tempfile
import threading
from collections import Counter
from datetime import date, datetime
from zoneinfo import ZoneInfo

import assets as asset_data
import fleet

log = logging.getLogger('asset_management.riga')

HERE = os.path.dirname(os.path.abspath(__file__))
MAPPING_PATH = os.path.join(HERE, 'riga_mapping.json')
# Riga's dates are plain calendar dates. This time zone is only used for the
# upload time, Riga's "today" and Riga's history lines.
RIGA_TIMEZONE = 'Europe/Riga'

# The fields read from each sheet. Only serial must have a header.
FIELDS = ('serial', 'condition', 'location', 'owner', 'notes', 'date')
# Text date formats the mapping may list, and what each means to Python.
DATE_TOKENS = {'dd.mm.yyyy': '%d.%m.%Y', 'dd.mm.yy': '%d.%m.%y'}
CONDITION_VALUES = ('FMC', 'PMC', 'NMC', 'Destroyed/Lost')

# A Riga row's ID is minus (sheet number x 10,000,000 + Excel row number).
# The sheet number is the sheet's place in riga_mapping.json, from 1. US IDs
# (SharePoint item IDs and CSV row positions) are always positive, and an
# Excel sheet has at most 1,048,576 rows, so no two IDs can ever be equal.
SHEET_ID_STEP = 10000000

# Uploads larger than this are refused. The sample workbook is far smaller.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_NAME_LENGTH = 80


class WorkbookError(Exception):
    """The file cannot be opened as an Excel workbook."""


class UploadRejected(Exception):
    """An upload was refused; the message is for the page."""


# ---------------------------------------------------------------------------
# The mapping file
# ---------------------------------------------------------------------------

def _norm(text):
    """Header and sheet names match ignoring case and surrounding spaces."""
    return ' '.join(str(text).split()).casefold() if text is not None else ''


def read_mapping(path=MAPPING_PATH):
    """Read riga_mapping.json. Raises OSError if it cannot be read and
    ValueError if it is not in the expected shape."""
    with open(path, encoding='utf-8') as f:
        try:
            data = json.load(f)
        except ValueError:
            raise ValueError('it is not valid JSON') from None
    if not isinstance(data, dict):
        raise ValueError('it does not hold a mapping')
    sheets = data.get('sheets')
    if not isinstance(sheets, list) or not sheets:
        raise ValueError('it has no "sheets" list')
    seen = set()
    for i, sheet in enumerate(sheets, start=1):
        if not isinstance(sheet, dict):
            raise ValueError('sheet entry %d is not a set of settings' % i)
        name = sheet.get('sheet')
        if not isinstance(name, str) or not name.strip():
            raise ValueError('sheet entry %d has no sheet name' % i)
        if _norm(name) in seen:
            raise ValueError('it names the sheet %s twice' % name)
        seen.add(_norm(name))
        if not isinstance(sheet.get('asset_type'), str) or not sheet['asset_type'].strip():
            raise ValueError('sheet %s has no asset_type' % name)
        if sheet.get('family') not in fleet.FAMILIES:
            raise ValueError('sheet %s has the family %r, which is not one of %s'
                             % (name, sheet.get('family'), ', '.join(fleet.FAMILIES)))
        row = sheet.get('header_row')
        if not isinstance(row, int) or isinstance(row, bool) or row < 1:
            raise ValueError('sheet %s needs a header_row of 1 or more' % name)
        columns = sheet.get('columns')
        if not isinstance(columns, dict) or set(columns) != set(FIELDS):
            raise ValueError('sheet %s needs "columns" with exactly %s' % (name, ', '.join(FIELDS)))
        for field, header in columns.items():
            if header is None and field != 'serial':
                continue
            if not isinstance(header, str) or not header.strip():
                raise ValueError('sheet %s has no header text for %s' % (name, field))
    ignored = data.get('ignored_sheets', [])
    if not isinstance(ignored, list) or not all(isinstance(s, str) for s in ignored):
        raise ValueError('"ignored_sheets" must be a list of sheet names')
    if seen & {_norm(s) for s in ignored}:
        raise ValueError('a sheet is both mapped and ignored')
    conditions = data.get('conditions')
    if not isinstance(conditions, dict) or not conditions:
        raise ValueError('it has no "conditions"')
    for value, condition in conditions.items():
        if condition not in CONDITION_VALUES:
            raise ValueError('the Status %s maps to %r, which is not one of %s'
                             % (value, condition, ', '.join(CONDITION_VALUES)))
    formats = data.get('date_formats')
    if not isinstance(formats, list) or not formats:
        raise ValueError('it has no "date_formats"')
    for fmt in formats:
        if fmt not in DATE_TOKENS:
            raise ValueError('the date format %r is not one of %s'
                             % (fmt, ', '.join(DATE_TOKENS)))
    return data


def cross_check(families, mapping):
    """asset_families.json decides families; each sheet's family in the
    mapping is a check on it. Returns (the families to use for Riga rows,
    [(sheet, Asset Type, mapping's family, file's family or None)]). A type
    whose two families disagree is left out, so its rows are Unmapped."""
    use = dict(families)
    mismatches = []
    for sheet in mapping['sheets']:
        in_file = families.get(sheet['asset_type'])
        if in_file != sheet['family']:
            mismatches.append((sheet['sheet'], sheet['asset_type'], sheet['family'], in_file))
            use.pop(sheet['asset_type'], None)
    return use, mismatches


def mismatch_text(mismatch):
    sheet, asset_type, riga_family, file_family = mismatch
    if file_family is None:
        where = 'asset_families.json does not name it'
    else:
        where = 'asset_families.json says %s' % file_family
    return ('Sheet %s (%s): riga_mapping.json says %s but %s, so its rows are Unmapped.'
            % (sheet, asset_type, riga_family, where))


# ---------------------------------------------------------------------------
# Reading cells
# ---------------------------------------------------------------------------

def _is_error(cell):
    from openpyxl.cell.cell import ERROR_CODES
    return getattr(cell, 'data_type', None) == 'e' or (
        isinstance(cell.value, str) and cell.value.strip() in ERROR_CODES)


def _text(value):
    """A cell as text. A whole number stays without a decimal point, so a
    serial stored as a number reads as it looks."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    return str(value).strip()


def _parse_date(value, formats):
    """(date or None, could not be read). Excel dates are taken as they are,
    with no time zone shift. Text is read day first, in the mapping's formats."""
    if value is None:
        return None, False
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None, False
        for fmt in formats:
            try:
                return datetime.strptime(text, DATE_TOKENS[fmt]).date(), False
            except ValueError:
                continue
    return None, True


def _cell(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------

def row_id(sheet_number, excel_row):
    return -(sheet_number * SHEET_ID_STEP + excel_row)


def _open(path):
    from openpyxl import load_workbook
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise WorkbookError('it could not be opened as an Excel workbook (%s)'
                            % type(e).__name__) from None


def read_workbook(path, mapping, tz=None):
    """Read every mapped sheet. Returns a dict:

    rows: the asset rows, in the shape assets._clean_row gives plus `site`
    found, mapped, ignored, unknown: sheet names in workbook order
    missing: mapped sheets the workbook does not have
    sheets: {sheet: {asset_type, family, rows, missing_headers, bad_dates,
             formula_errors, has_date, samples}}
    unknown_conditions: {Status value: rows}
    dated_types: Asset Types whose sheet has a date column that was found

    Raises WorkbookError if the file is not an Excel workbook."""
    tz = tz or ZoneInfo(RIGA_TIMEZONE)
    formats = mapping['date_formats']
    conditions = {k.strip(): v for k, v in mapping['conditions'].items()}
    ignored = {_norm(s) for s in mapping.get('ignored_sheets', [])}
    by_name = {_norm(s['sheet']): (n, s) for n, s in enumerate(mapping['sheets'], start=1)}
    wb = _open(path)
    try:
        result = {'rows': [], 'found': list(wb.sheetnames), 'mapped': [], 'ignored': [],
                  'unknown': [], 'missing': [], 'sheets': {}, 'unknown_conditions': Counter(),
                  'dated_types': set()}
        present = set()
        for name in wb.sheetnames:
            key = _norm(name)
            if key in by_name:
                result['mapped'].append(name)
                present.add(key)
                number, sheet = by_name[key]
                result['sheets'][sheet['sheet']] = _read_sheet(
                    wb[name], number, sheet, conditions, formats, tz, result)
            elif key in ignored:
                result['ignored'].append(name)
            else:
                result['unknown'].append(name)
        result['missing'] = [s['sheet'] for s in mapping['sheets'] if _norm(s['sheet']) not in present]
    finally:
        wb.close()
    result['unknown_conditions'] = dict(result['unknown_conditions'].most_common())
    return result


def _read_sheet(ws, number, sheet, conditions, formats, tz, result):
    stats = {'asset_type': sheet['asset_type'], 'family': sheet['family'], 'rows': 0,
             'missing_headers': [], 'bad_dates': 0, 'formula_errors': 0,
             'has_date': False, 'samples': []}
    header_row = sheet['header_row']
    headers = {}
    for cells in ws.iter_rows(min_row=header_row, max_row=header_row):
        for i, cell in enumerate(cells):
            key = _norm(cell.value)
            if key and key not in headers:
                headers[key] = i
    where = {}
    for field in FIELDS:
        header = sheet['columns'][field]
        if header is None:
            where[field] = None
            continue
        where[field] = headers.get(_norm(header))
        if where[field] is None:
            stats['missing_headers'].append((field, header))
    stats['has_date'] = where['date'] is not None
    if stats['has_date']:
        result['dated_types'].add(sheet['asset_type'])
    if where['serial'] is None:
        return stats

    for excel_row, cells in enumerate(ws.iter_rows(min_row=header_row + 1), start=header_row + 1):
        serial_cell = _cell(cells, where['serial'])
        if serial_cell is None or serial_cell.value is None:
            continue
        if _is_error(serial_cell):
            stats['formula_errors'] += 1
            continue
        serial = _text(serial_cell.value)
        if not serial:
            continue
        values = {}
        for field in ('condition', 'location', 'owner', 'notes', 'date'):
            cell = _cell(cells, where[field])
            if cell is not None and _is_error(cell):
                stats['formula_errors'] += 1
                values[field] = None
            else:
                values[field] = cell.value if cell is not None else None
        status = _text(values['condition'])
        condition = conditions.get(status, '') if status else ''
        if status and not condition:
            result['unknown_conditions'][status] += 1
        when, bad = _parse_date(values['date'], formats)
        stats['bad_dates'] += bad
        location = _text(values['location'])
        row = asset_data._clean_row({
            'Serial Number': serial,
            'Asset Type': sheet['asset_type'],
            'Ownership': _text(values['owner']),
            'Location': [location] if location else [],
            'Material Condition': condition,
            'Asset Notes': _text(values['notes']),
        }, tz, row_id(number, excel_row))
        # Plain calendar dates: set here, never shifted by a time zone.
        row['last_inventory'] = when
        result['rows'].append(asset_data.with_site(row, fleet.RIGA))
        stats['rows'] += 1
        if len(stats['samples']) < 3:
            stats['samples'].append({'excel_row': excel_row, 'serial': row['serial_number'],
                                     'status': status, 'condition': condition,
                                     'location': location, 'owner': row['ownership'],
                                     'date': when, 'notes_length': len(row['notes'])})
    return stats


def findings(result):
    """The problems found in a workbook, as sentences for the log and the
    page. They name sheets, headers and Status values, never row contents."""
    out = []
    for name in result['unknown']:
        out.append('Sheet %s is not in riga_mapping.json and was skipped.' % name)
    for name in result['missing']:
        out.append('Sheet %s is in riga_mapping.json but not in the workbook.' % name)
    for name, s in result['sheets'].items():
        for field, header in s['missing_headers']:
            out.append('Sheet %s has no column headed "%s", so its %s is blank.'
                       % (name, header, field))
        if s['bad_dates']:
            out.append('Sheet %s: %d date(s) could not be read and are blank.'
                       % (name, s['bad_dates']))
        if s['formula_errors']:
            out.append('Sheet %s: %d formula error(s), such as #REF!, in the columns read; '
                       'those cells are blank.' % (name, s['formula_errors']))
    for value, n in result['unknown_conditions'].items():
        out.append('Status "%s" on %d row(s) is not in riga_mapping.json, so those rows have '
                   'no condition.' % (value, n))
    return out


# ---------------------------------------------------------------------------
# Alerts
# ---------------------------------------------------------------------------

def attention(assets, today, stale_days, dated_types):
    """The Riga alerts, in the same shape as assets.attention. Riga has no
    checkout data, so only these four rules apply. The two inventory rules
    only look at sheets with a date column; that date is when the row was
    last updated and signed, not a physical count."""
    rules = [
        ('no_condition', 'No condition',
         'Status is blank, or holds a value riga_mapping.json does not know.',
         lambda a: not a['condition']),
        ('never_inventoried', 'Never inventoried',
         'The last-updated date is blank (sheets with a date column only).',
         lambda a: a['asset_type'] in dated_types and a['last_inventory'] is None),
        ('inventory_overdue', 'Inventory overdue',
         'The last-updated date is more than %d days ago (sheets with a date column only).'
         % stale_days,
         lambda a: a['asset_type'] in dated_types and a['last_inventory'] is not None
         and (today - a['last_inventory']).days > stale_days),
        ('degraded_no_notes', 'PMC/NMC with no notes',
         'Status means PMC or NMC but the notes column says nothing about why.',
         lambda a: a['condition'] in ('PMC', 'NMC') and not a['notes']),
    ]
    return [{'key': key, 'title': title, 'why': why, 'assets': [a for a in assets if test(a)]}
            for key, title, why, test in rules]


# ---------------------------------------------------------------------------
# The data source the app uses
# ---------------------------------------------------------------------------

def upload_info_path(path):
    """The small JSON file beside the workbook that says who uploaded it."""
    return os.path.splitext(path)[0] + '.upload.json'


class RigaSource:
    """The uploaded workbook at `path`, read again whenever it or the mapping
    file changes. load() never raises: a problem is logged once, and again
    only when it changes, and the Riga view shows a short note."""

    def __init__(self, path, mapping_path=MAPPING_PATH, off_reason=None):
        self.path = path
        self.mapping_path = mapping_path
        # Set when the folder could not be made at startup.
        self.off_reason = off_reason
        self.tz = ZoneInfo(RIGA_TIMEZONE)
        self._lock = threading.Lock()
        self._cache_key = None
        self._cache = None
        self._logged = None
        self.upload_lock = threading.Lock()

    def _log_once(self, key, level, message, *args):
        if key != self._logged:
            self._logged = key
            log.log(level, message, *args)

    def mapping(self):
        """(mapping, problem): problem is a sentence for the log, or None."""
        name = os.path.basename(self.mapping_path)
        try:
            return read_mapping(self.mapping_path), None
        except FileNotFoundError:
            return None, 'there is no %s next to the app' % name
        except OSError as e:
            return None, '%s could not be opened (%s)' % (name, type(e).__name__)
        except ValueError as e:
            return None, '%s could not be used: %s' % (name, e)

    def upload_info(self):
        """Who uploaded the current workbook, when, and from which file, or
        None if that was not recorded."""
        try:
            with open(upload_info_path(self.path), encoding='utf-8') as f:
                info = json.load(f)
            info['uploaded_at'] = datetime.fromisoformat(info['uploaded_at']).astimezone(self.tz)
            return info
        except FileNotFoundError:
            return None
        except (OSError, ValueError, KeyError, TypeError) as e:
            log.warning('Riga workbook: the upload record %s could not be read (%s).',
                        upload_info_path(self.path), type(e).__name__)
            return None

    def load(self):
        """{'state': 'off' | 'mapping' | 'none' | 'unreadable' | 'ok', ...}.
        With 'ok' it also holds 'data' (the rows, in the shape the pages use),
        'result' (from read_workbook), 'mapping' and 'findings'."""
        if self.off_reason:
            return {'state': 'off', 'problem': self.off_reason}
        mapping, problem = self.mapping()
        if problem:
            self._log_once(('mapping', problem), logging.ERROR,
                           'Riga view unavailable, because %s. The US view is not affected.',
                           problem)
            return {'state': 'mapping', 'problem': problem}
        try:
            st = os.stat(self.path)
        except FileNotFoundError:
            self._log_once(('none',), logging.INFO,
                           'Riga workbook: none uploaded yet (%s).', self.path)
            return {'state': 'none'}
        except OSError as e:
            problem = '%s could not be read (%s: %s)' % (self.path, type(e).__name__, e)
            self._log_once(('unreadable', problem), logging.ERROR, 'Riga workbook: %s.', problem)
            return {'state': 'unreadable', 'problem': problem}
        key = (st.st_mtime_ns, st.st_size, json.dumps(mapping, sort_keys=True))
        with self._lock:
            if key == self._cache_key:
                cached = dict(self._cache)
            else:
                try:
                    result = read_workbook(self.path, mapping, self.tz)
                except WorkbookError as e:
                    problem = '%s: %s' % (self.path, e)
                    self._log_once(('unreadable', problem), logging.ERROR,
                                   'Riga workbook: %s. The US view is not affected.', problem)
                    return {'state': 'unreadable', 'problem': problem}
                found = findings(result)
                log.info('Riga workbook: %d rows read from %s.', len(result['rows']), self.path)
                for line in found:
                    log.warning('Riga workbook: %s', line)
                self._logged = ('ok',)
                self._cache_key = key
                self._cache = {'state': 'ok', 'result': result, 'mapping': mapping,
                               'findings': found, 'mtime': st.st_mtime}
                cached = dict(self._cache)
        info = self.upload_info()
        loaded_at = info['uploaded_at'] if info else datetime.fromtimestamp(cached['mtime'], self.tz)
        name = info.get('original_filename') if info else None
        cached['upload'] = info
        cached['data'] = {
            'fields': {}, 'assets': cached['result']['rows'], 'loaded_at': loaded_at,
            'source': 'Riga workbook' + (' %s' % name if name else ''), 'live': False,
            'stale': None,
        }
        return cached

    def forget(self):
        with self._lock:
            self._cache_key = None
            self._cache = None


# ---------------------------------------------------------------------------
# Uploading
# ---------------------------------------------------------------------------

def _stamp(now):
    return '%04d%02d%02d-%02d%02d%02d' % (now.year, now.month, now.day,
                                          now.hour, now.minute, now.second)


def backup_path(path, now):
    """current.xlsx -> current-20261029-143012.xlsx, with -2, -3 ... added if
    that name is already taken."""
    stem, ext = os.path.splitext(path)
    candidate = '%s-%s%s' % (stem, _stamp(now), ext)
    n = 2
    while os.path.exists(candidate):
        candidate = '%s-%s-%d%s' % (stem, _stamp(now), n, ext)
        n += 1
    return candidate


def _write_json(path, data):
    folder = os.path.dirname(path)
    fd, tmp = tempfile.mkstemp(dir=folder, prefix='.upload-', suffix='.json.part')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8', newline='\n') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.write('\n')
        os.replace(tmp, path)
    finally:
        _remove_quietly(tmp)


def _remove_quietly(path):
    """Remove a leftover temporary file. A failure is logged, never raised,
    so it cannot hide the real result of an upload."""
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError as e:
        log.warning('Riga upload: the temporary file %s could not be removed (%s).',
                    path, type(e).__name__)


def save_upload(source, stream, filename, name, now, limit=None):
    """Check an uploaded workbook and put it in place. Returns the upload
    record. Raises UploadRejected, with a reason for the page, and then the
    previous workbook is untouched.

    Steps: write to a temporary file in the same folder, check it, keep the
    previous workbook as a dated backup, move the new one into place, then
    write the upload record beside it."""
    limit = MAX_UPLOAD_BYTES if limit is None else limit
    if source.off_reason:
        raise UploadRejected('The folder for the Riga workbook could not be created when the '
                             'dashboard started, so nothing can be saved. The reason is in the '
                             'dashboard log.')
    name = (name or '').strip()
    if not name:
        raise UploadRejected('Enter your name, so the Riga view can say who uploaded the file.')
    if len(name) > MAX_NAME_LENGTH:
        raise UploadRejected('Your name can be at most %d characters.' % MAX_NAME_LENGTH)
    original = re.split(r'[\\/]', filename or '')[-1].strip()
    if not original:
        raise UploadRejected('Choose a file to upload.')
    if not original.lower().endswith('.xlsx'):
        raise UploadRejected('Only Excel workbooks ending in .xlsx can be uploaded. '
                             'This file is %s.' % original)
    mapping, problem = source.mapping()
    if problem:
        log.error('Riga upload refused, because %s.', problem)
        raise UploadRejected('The file cannot be checked, because the Riga mapping file '
                             '(riga_mapping.json) could not be read. The reason is in the '
                             'dashboard log.')
    folder = os.path.dirname(source.path)
    with source.upload_lock:
        # openpyxl only opens names ending in .xlsx.
        fd, tmp = tempfile.mkstemp(dir=folder, prefix='.upload-', suffix='.part.xlsx')
        try:
            size = 0
            with os.fdopen(fd, 'wb') as out:
                while True:
                    chunk = stream.read(64 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > limit:
                        raise UploadRejected('The file is larger than the %s limit.'
                                             % size_text(limit))
                    out.write(chunk)
            try:
                result = read_workbook(tmp, mapping, source.tz)
            except WorkbookError:
                raise UploadRejected('%s could not be opened as an Excel workbook. Save it from '
                                     'Excel as an .xlsx workbook and try again.' % original) from None
            if not result['rows']:
                raise UploadRejected('%s holds no Riga assets: none of the sheets named in '
                                     'riga_mapping.json has a row with a serial number.' % original)
            backup = None
            if os.path.exists(source.path):
                backup = backup_path(source.path, now)
                shutil.copy2(source.path, backup)
            try:
                os.replace(tmp, source.path)
            except OSError as e:
                log.error('Riga upload: %s could not be replaced (%s: %s).',
                          source.path, type(e).__name__, e)
                raise UploadRejected('The new file could not be put in place (%s). The previous '
                                     'workbook is still in use. If the workbook is open on the '
                                     'server, close it and try again.' % type(e).__name__) from None
            info = {
                'uploaded_by': name,
                'uploaded_at': now.isoformat(timespec='seconds'),
                'original_filename': original,
                'size_bytes': size,
                'rows_per_sheet': {s: v['rows'] for s, v in result['sheets'].items()},
            }
            _write_json(upload_info_path(source.path), info)
        finally:
            _remove_quietly(tmp)
    source.forget()
    log.info('Riga workbook uploaded by %s: %s, %d rows%s.', name, original, len(result['rows']),
             ', previous file kept as %s' % os.path.basename(backup) if backup else '')
    info['backup'] = backup
    info['rows'] = len(result['rows'])
    return info


def size_text(n):
    if n >= 1024 * 1024:
        return '%g MB' % round(n / (1024 * 1024), 1)
    if n >= 1024:
        return '%g KB' % round(n / 1024, 1)
    return '%d bytes' % n


# ---------------------------------------------------------------------------
# The check command
# ---------------------------------------------------------------------------

def mask(serial):
    """The first characters of a serial number, then ***. At most 3, and
    never more than half, so a short serial is not shown whole."""
    serial = serial or ''
    return serial[:min(3, len(serial) // 2)] + '***'


def _names(names):
    return ', '.join(names) if names else 'none'


def check(path, mapping_path=MAPPING_PATH, families_path=fleet.ASSET_FAMILIES_PATH, out=None):
    """Print what the loader makes of a workbook. Returns the exit code."""
    out = out or sys.stdout

    def say(text=''):
        out.write(text + '\n')

    say('Riga workbook check: %s' % os.path.basename(path))
    try:
        mapping = read_mapping(mapping_path)
    except (OSError, ValueError) as e:
        say('FAILED: %s could not be used: %s' % (os.path.basename(mapping_path), e))
        return 1
    say('Mapping: %s, %d sheets' % (os.path.basename(mapping_path), len(mapping['sheets'])))
    try:
        result = read_workbook(path, mapping)
    except WorkbookError as e:
        say('FAILED: the file %s.' % e)
        return 1
    except OSError as e:
        say('FAILED: the file could not be read (%s).' % type(e).__name__)
        return 1

    say()
    say('1. Sheets')
    say('   Found (%d): %s' % (len(result['found']), _names(result['found'])))
    say('   Mapped (%d): %s' % (len(result['mapped']), _names(result['mapped'])))
    say('   Ignored (%d): %s' % (len(result['ignored']), _names(result['ignored'])))
    say('   Unknown, skipped (%d): %s' % (len(result['unknown']), _names(result['unknown'])))
    say('   In the mapping but not in the workbook (%d): %s'
        % (len(result['missing']), _names(result['missing'])))

    say()
    say('2. Rows read per sheet')
    for name, s in result['sheets'].items():
        say('   %s (%s, %s): %d' % (name, s['asset_type'], s['family'], s['rows']))
    say('   Total: %d' % len(result['rows']))

    say()
    say('3. Mapped headers not found')
    missing = [(n, f, h) for n, s in result['sheets'].items() for f, h in s['missing_headers']]
    for name, field, header in missing:
        say('   %s: "%s" (%s)' % (name, header, field))
    if not missing:
        say('   none')

    say()
    say('4. Unknown Status values (read as no condition)')
    for value, n in result['unknown_conditions'].items():
        say('   "%s": %d row(s)' % (value, n))
    if not result['unknown_conditions']:
        say('   none')

    say()
    say('5. Dates that could not be read, and formula errors, per sheet')
    any_problem = False
    for name, s in result['sheets'].items():
        if s['bad_dates'] or s['formula_errors']:
            any_problem = True
            say('   %s: %d date(s) not read, %d formula error(s)'
                % (name, s['bad_dates'], s['formula_errors']))
        if not s['has_date']:
            say('   %s: no date column, so the inventory alerts skip it' % name)
    if not any_problem:
        say('   no unreadable dates or formula errors')

    say()
    say('6. First 3 rows of each sheet (serials masked, notes shown as their length)')
    for name, s in result['sheets'].items():
        say('   %s' % name)
        for r in s['samples']:
            say('     row %d: serial %s | Status %s -> %s | Location %s | Owner %s | date %s | '
                'notes %d characters'
                % (r['excel_row'], mask(r['serial']), r['status'] or '(blank)',
                   r['condition'] or 'no condition', r['location'] or '(blank)',
                   r['owner'] or '(blank)', r['date'].isoformat() if r['date'] else '(blank)',
                   r['notes_length']))
        if not s['samples']:
            say('     no rows')

    say()
    say('7. Totals per family')
    try:
        families = fleet.read_families(families_path)
    except (OSError, ValueError) as e:
        say('   %s could not be used (%s), so every row is Unmapped.'
            % (os.path.basename(families_path), e))
        families = {}
    use, mismatches = cross_check(families, mapping)
    ready = fleet.fleet_readiness(result['rows'], use)
    for family, r in ready['families'].items():
        if r['total']:
            say('   %s: %d (Ready %d, Limited %d, Not Ready %d, No condition %d, '
                'Destroyed/Lost %d)' % (family, r['total'], r['fmc'], r['pmc'], r['nmc'],
                                       r['none'], r['lost']))
    for m in mismatches:
        say('   WARNING: %s' % mismatch_text(m))
    pct = ready['overall']['pct']
    say('   Headline readiness, engines left out: %s'
        % ('none' if pct is None else '%.0f%%' % pct))
    return 0


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2 or argv[0] != '--check':
        sys.stderr.write('Usage: python -m riga_loader --check PATH\n')
        return 2
    return check(argv[1])


if __name__ == '__main__':
    sys.exit(main())
